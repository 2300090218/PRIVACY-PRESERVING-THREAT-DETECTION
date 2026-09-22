"""
Script to generate a professional Word (.docx) document containing all
backend API keys, agent credentials, and user JWT authentication tokens.
"""

import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_credentials_doc():
    doc = Document()

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Styles
    navy = RGBColor(16, 44, 87)
    slate = RGBColor(51, 65, 85)
    dark_gray = RGBColor(71, 85, 105)
    code_blue = RGBColor(30, 58, 138)

    # Document Header Title
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = title.add_run("Privacy-Preserving Threat Detection Platform")
    run_title.font.name = "Arial"
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = navy

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = subtitle.add_run("Backend API Keys, Agent Credentials & Authentication Reference")
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(13)
    run_sub.font.color.rgb = slate
    run_sub.italic = True

    doc.add_paragraph() # Spacing

    # Section 1: Overview Callout Box
    table_callout = doc.add_table(rows=1, cols=1)
    table_callout.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell_callout = table_callout.rows[0].cells[0]
    set_cell_background(cell_callout, "F1F5F9")
    set_cell_margins(cell_callout, top=140, bottom=140, left=200, right=200)

    p_callout = cell_callout.paragraphs[0]
    r_c1 = p_callout.add_run("SECURITY POLICY NOTICE: ")
    r_c1.bold = True
    r_c1.font.size = Pt(10)
    r_c1.font.color.rgb = navy
    r_c2 = p_callout.add_run(
        "All local telemetry must be transformed by the edge Privacy Gateway before reaching the central server. "
        "Every edge agent must authenticate with an authorized X-API-Key and X-Organization-ID. "
        "User sessions on the SOC dashboard require OAuth2 Bearer JWT tokens issued by /api/v1/auth/login."
    )
    r_c2.font.size = Pt(10)
    r_c2.font.color.rgb = slate

    doc.add_paragraph()

    # Section 2: Edge Agent API Keys
    h1 = doc.add_paragraph()
    r_h1 = h1.add_run("1. Edge Agent API Keys (Local Organization Ingestion)")
    r_h1.font.name = "Arial"
    r_h1.font.size = Pt(14)
    r_h1.font.bold = True
    r_h1.font.color.rgb = navy

    desc1 = doc.add_paragraph()
    r_d1 = desc1.add_run(
        "These API keys are used by local organization edge agents to authenticate telemetry transmission "
        "to the Central Ingestion API (POST /api/v1/events). Prohibited fields must be minimized prior to dispatch."
    )
    r_d1.font.size = Pt(10.5)
    r_d1.font.color.rgb = dark_gray

    # Table of Agent Keys
    agent_data = [
        ("Organization", "Agent ID", "Agent Name", "Active API Key", "Header"),
        ("Organization A\n(org_enterprise_a)", "agent-dmz-01", "DMZ Gateway Edge Agent", "agent_key_enterprise_a_dmz_prod_secret", "X-API-Key"),
        ("Organization A\n(org_enterprise_a)", "agent-sec-alpha", "Alpha Security Sensor", "agkey_8d1583c1e86c4530b43cc9be24122461", "X-API-Key"),
        ("Organization B\n(org_finance_b)", "agent-sec-beta", "Finance Security Sensor", "agkey_3edcaf55c08f41739eaeeccd2fab8532", "X-API-Key"),
        ("Organization C\n(org_cloud_c)", "agent-sec-gamma", "Cloud VPC Security Sensor", "agkey_c32c5ff27c7241c38f04d95997a56dc0", "X-API-Key"),
    ]

    table_agents = doc.add_table(rows=len(agent_data), cols=5)
    table_agents.alignment = WD_TABLE_ALIGNMENT.CENTER

    for row_idx, row_values in enumerate(agent_data):
        row = table_agents.rows[row_idx]
        is_header = (row_idx == 0)
        for col_idx, text in enumerate(row_values):
            cell = row.cells[col_idx]
            cell.text = text
            set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
            if is_header:
                set_cell_background(cell, "1E3A8A")
                for p in cell.paragraphs:
                    for r in p.runs:
                        r.font.bold = True
                        r.font.size = Pt(9.5)
                        r.font.color.rgb = RGBColor(255, 255, 255)
            else:
                bg = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
                set_cell_background(cell, bg)
                for p in cell.paragraphs:
                    for r in p.runs:
                        r.font.size = Pt(9)
                        if col_idx == 3:
                            r.font.name = "Consolas"
                            r.font.color.rgb = code_blue

    doc.add_paragraph()

    # Section 3: User Accounts & JWT Tokens
    h2 = doc.add_paragraph()
    r_h2 = h2.add_run("2. User Accounts & JWT Bearer Tokens (SOC Management)")
    r_h2.font.name = "Arial"
    r_h2.font.size = Pt(14)
    r_h2.font.bold = True
    r_h2.font.color.rgb = navy

    desc2 = doc.add_paragraph()
    r_d2 = desc2.add_run(
        "These credentials grant access to administrative routes, live alerts, detections, and the Next.js SOC Dashboard. "
        "Authentication uses standard OAuth2 Bearer Tokens (Authorization: Bearer <token>)."
    )
    r_d2.font.size = Pt(10.5)
    r_d2.font.color.rgb = dark_gray

    # User Accounts Details
    users_info = [
        {
            "role": "Platform Administrator",
            "username": "admin",
            "password": "AdminPass123!",
            "org": "org_enterprise_a",
            "token": "eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9.eyJzdWIiOiJhZG1pbiIsInJvbGUiOiJBRE1JTiIsIm9yZ2FuaXphdGlvbl9pZCI6Im9yZ19lbnRlcnByaXNlX2EiLCJleHAiOjE3OTAxNDU1NjJ9.ZJRJdlf6uiEqobQbKOe54r9Xz-g-cMNGaaYI1dEur_k"
        },
        {
            "role": "Security Analyst",
            "username": "analyst",
            "password": "AnalystPass123!",
            "org": "org_enterprise_a",
            "token": "eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9.eyJzdWIiOiJhbmFseXN0Iiwicm9sZSI6IlNFQ1VSSVRZX0FOQUxZU1QiLCJvcmdhbml6YXRpb25faWQiOiJvcmdfZW50ZXJwcmlzZV9hIiwiZXhwIjoxNzkwMTQ1NTYyfQ.ooK2xy6CvTtr05pYi24Y7pTHsCvMXwi8cLHqi_OtIZY"
        }
    ]

    for u in users_info:
        p_u = doc.add_paragraph()
        r_ut = p_u.add_run(f"• {u['role']} ({u['username']})")
        r_ut.bold = True
        r_ut.font.size = Pt(11)
        r_ut.font.color.rgb = navy

        table_u = doc.add_table(rows=4, cols=2)
        table_u.alignment = WD_TABLE_ALIGNMENT.CENTER
        rows_u = [
            ("Username / Password:", f"{u['username']}  /  {u['password']}"),
            ("Organization ID:", u['org']),
            ("Assigned RBAC Role:", u['role']),
            ("Active JWT Token:", u['token'])
        ]
        for idx, (label, val) in enumerate(rows_u):
            r = table_u.rows[idx]
            r.cells[0].text = label
            r.cells[1].text = val
            set_cell_background(r.cells[0], "F1F5F9")
            set_cell_background(r.cells[1], "FFFFFF")
            set_cell_margins(r.cells[0], top=80, bottom=80, left=120, right=120)
            set_cell_margins(r.cells[1], top=80, bottom=80, left=120, right=120)
            r.cells[0].paragraphs[0].runs[0].font.bold = True
            r.cells[0].paragraphs[0].runs[0].font.size = Pt(9.5)
            p_val = r.cells[1].paragraphs[0].runs[0]
            p_val.font.size = Pt(8.5)
            if idx == 3:
                p_val.font.name = "Consolas"
                p_val.font.color.rgb = code_blue

        doc.add_paragraph()

    # Section 4: Environment Variables & Backend Secrets
    h3 = doc.add_paragraph()
    r_h3 = h3.add_run("3. System Environment Secrets (.env)")
    r_h3.font.name = "Arial"
    r_h3.font.size = Pt(14)
    r_h3.font.bold = True
    r_h3.font.color.rgb = navy

    env_data = [
        ("Variable Name", "Current Configured Value", "Purpose"),
        ("SECRET_KEY", "threat-detection-production-secret-key-32-chars-min!", "HMAC secret for signing JWT analyst/admin tokens"),
        ("ALGORITHM", "HS256", "Cryptographic signature algorithm for access tokens"),
        ("ACCESS_TOKEN_EXPIRE_MINUTES", "1440 (24 hours)", "JWT token validity window"),
        ("DATABASE_URL", "postgresql+asyncpg://threat_user:threat_pass@localhost:5432/threat_detection", "Async connection string (SQLite dev fallback supported)"),
        ("CENTRAL_API_URL", "http://127.0.0.1:8000", "Base endpoint for agent telemetry dispatch"),
    ]

    table_env = doc.add_table(rows=len(env_data), cols=3)
    table_env.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row_idx, row_values in enumerate(env_data):
        row = table_env.rows[row_idx]
        is_header = (row_idx == 0)
        for col_idx, text in enumerate(row_values):
            cell = row.cells[col_idx]
            cell.text = text
            set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
            if is_header:
                set_cell_background(cell, "1E3A8A")
                for p in cell.paragraphs:
                    for r in p.runs:
                        r.font.bold = True
                        r.font.size = Pt(9.5)
                        r.font.color.rgb = RGBColor(255, 255, 255)
            else:
                set_cell_background(cell, "F8FAFC" if row_idx % 2 == 1 else "FFFFFF")
                for p in cell.paragraphs:
                    for r in p.runs:
                        r.font.size = Pt(9)
                        if col_idx == 0:
                            r.font.name = "Consolas"
                            r.font.bold = True

    doc.add_paragraph()

    # Section 5: HTTP Usage Example
    h4 = doc.add_paragraph()
    r_h4 = h4.add_run("4. HTTP Request Header Formats")
    r_h4.font.name = "Arial"
    r_h4.font.size = Pt(14)
    r_h4.font.bold = True
    r_h4.font.color.rgb = navy

    p_code1 = doc.add_paragraph()
    r_c_title1 = p_code1.add_run("A. Edge Telemetry Ingestion (Agent API Key Header):\n")
    r_c_title1.bold = True
    r_c_title1.font.size = Pt(10)
    r_code1 = p_code1.add_run(
        "POST /api/v1/events HTTP/1.1\n"
        "Host: 127.0.0.1:8000\n"
        "Content-Type: application/json\n"
        "X-Organization-ID: org_enterprise_a\n"
        "X-API-Key: agkey_8d1583c1e86c4530b43cc9be24122461"
    )
    r_code1.font.name = "Consolas"
    r_code1.font.size = Pt(9)
    r_code1.font.color.rgb = code_blue

    p_code2 = doc.add_paragraph()
    r_c_title2 = p_code2.add_run("B. REST API Management (JWT Bearer Token Header):\n")
    r_c_title2.bold = True
    r_c_title2.font.size = Pt(10)
    r_code2 = p_code2.add_run(
        "GET /api/v1/alerts HTTP/1.1\n"
        "Host: 127.0.0.1:8000\n"
        "Authorization: Bearer eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9...\n"
        "Accept: application/json"
    )
    r_code2.font.name = "Consolas"
    r_code2.font.size = Pt(9)
    r_code2.font.color.rgb = code_blue

    # Save document
    output_path = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "API_KEYS_AND_CREDENTIALS.docx"))
    doc.save(output_path)
    print(f"[SUCCESS] Word document generated at: {output_path}")
    return output_path

if __name__ == "__main__":
    create_credentials_doc()
